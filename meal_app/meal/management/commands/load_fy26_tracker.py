import re
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from meal_app.meal.models import Indicator, Output, QuarterPlan, StrategicObjective


class Command(BaseCommand):
    help = 'Load the DDS-NK FY26 M&E framework and people/non-people activity trackers from an .xlsx workbook.'

    def add_arguments(self, parser):
        parser.add_argument('workbook', nargs='?', default='Activity Tracker FY26.xlsx')

    def handle(self, *args, **options):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise CommandError('openpyxl is required to load the FY26 workbook.') from exc
        path = Path(options['workbook'])
        if not path.exists():
            raise CommandError(f'Workbook not found: {path}')
        wb = load_workbook(path, data_only=False, read_only=True)
        if 'IPTT' in wb.sheetnames:
            cached_wb = load_workbook(path, data_only=True, read_only=True)
            count = self._load_iptt(cached_wb['IPTT'])
            cached_wb.close()
            wb.close()
            self.stdout.write(self.style.SUCCESS(f'Loaded {count} FY26 IPTT activities and quarter plans.'))
            return
        people = self._tracker_names(wb['People Activity Tracker'], 2)
        non_people = self._tracker_names(wb['Non People Activity Tracker'], 1)
        objective = output = None
        count = 0
        for row in wb['M&E Framework'].iter_rows(min_row=3, values_only=True):
            first, second = self._text(row[0]), self._text(row[1])
            if first.startswith('Outcome '):
                code, title = self._split(first)
                objective, _ = StrategicObjective.objects.update_or_create(code=code, defaults={'title': title})
                output = None
                continue
            if first.startswith('Output '):
                if not objective: continue
                code, title = self._split(first)
                output, _ = Output.objects.update_or_create(code=code, defaults={'title': title, 'objective': objective})
                continue
            if not output or not first.startswith('Activity '): continue
            match = re.search(r'Activity\s+([0-9]+(?:\.[0-9]+)+)', first)
            if not match: continue
            code = match.group(1)
            values = list(row) + [None] * 11
            kind = 'people' if first in people or code in people else 'non_people' if first in non_people or code in non_people else 'people'
            Indicator.objects.update_or_create(code=code, defaults={
                'title': re.sub(r'^Activity\s+[0-9.]+:\s*', '', first).strip() or first,
                'output': output, 'activity_type': kind, 'indicator_text': second,
                'indicator_type': self._text(values[2]), 'numerator': self._text(values[3]), 'denominator': self._text(values[4]),
                'baseline': self._number(values[5]), 'target': self._number(values[6]), 'data_source': self._text(values[7]),
                'reporting_frequency': self._text(values[8]), 'responsible': self._text(values[9]), 'notes': self._text(values[10]),
            })
            count += 1
        wb.close()
        self.stdout.write(self.style.SUCCESS(f'Loaded {count} FY26 activities.'))

    def _load_iptt(self, sheet):
        """Load the attached workbook's IPTT sheet into the MEAL model."""
        current_objective = None
        current_output = None
        count = 0
        quarter_columns = {
            "Q1": (3, 7, 8, 15),
            "Q2": (4, 9, 10, 16),
            "Q3": (5, 11, 12, 17),
            "Q4": (6, 13, 14, 18),
        }

        for row_number, row in enumerate(sheet.iter_rows(min_row=3, max_col=24, values_only=True), start=3):
            values = list(row) + [None] * 24
            hierarchy = self._text(values[1])
            indicator_value = self._text(values[2])

            if hierarchy.lower().startswith('outcome '):
                code, title = self._split_label(hierarchy, 'Outcome')
                current_objective, _ = StrategicObjective.objects.update_or_create(
                    code=code,
                    defaults={'title': title},
                )
                current_output = None
                continue

            if hierarchy.lower().startswith(('output ', 'ouput ')) or hierarchy.startswith('4.'):
                if hierarchy.startswith('4.'):
                    code, title = '4', hierarchy[2:].lstrip(' .:')
                else:
                    code, title = self._split_label(hierarchy, 'Output')
                if not current_objective:
                    current_objective, _ = StrategicObjective.objects.get_or_create(
                        code=code.split('.')[0],
                        defaults={'title': 'FY26 programme outputs'},
                    )
                current_output, _ = Output.objects.update_or_create(
                    code=code,
                    defaults={'title': title, 'objective': current_objective},
                )

            if not indicator_value or ':' not in indicator_value:
                continue

            indicator_code, indicator_title = indicator_value.split(':', 1)
            indicator_code = indicator_code.strip()
            indicator_title = indicator_title.strip()
            if not indicator_code or not current_output:
                continue

            annual_target = self._number(values[20])
            indicator, _ = Indicator.objects.update_or_create(
                code=indicator_code,
                defaults={
                    'title': indicator_title or indicator_code,
                    'output': current_output,
                    'indicator_text': indicator_title,
                    'target': annual_target,
                    'active': True,
                },
            )

            for quarter, (target_col, planned_col, expense_col, result_col) in quarter_columns.items():
                QuarterPlan.objects.update_or_create(
                    indicator=indicator,
                    reporting_year=2026,
                    quarter=quarter,
                    defaults={
                        'target': self._number(values[target_col]),
                        'planned_budget': self._number(values[planned_col]) or Decimal('0'),
                        'budget_expense': self._number(values[expense_col]) or Decimal('0'),
                        'workbook_result': self._number(values[result_col]),
                        'source_row': row_number,
                    },
                )
            count += 1
        return count

    @staticmethod
    def _text(value): return str(value).strip() if value not in (None, '') else ''

    @staticmethod
    def _number(value):
        if value in (None, '', 'N/A') or (isinstance(value, str) and value.startswith('=')): return None
        try: return Decimal(str(value).replace(',', '').strip())
        except (InvalidOperation, ValueError): return None

    @staticmethod
    def _split(value):
        match = re.match(r'^(?:Outcome|Output)\s+([0-9.]+):?\s*(.*)$', value)
        return (match.group(1), match.group(2).strip()) if match else (value, value)

    @staticmethod
    def _split_label(value, prefix):
        prefix_pattern = r'Ou(?:t)?put' if prefix.lower() == 'output' else re.escape(prefix)
        match = re.match(rf'^{prefix_pattern}\s+([0-9A-Za-z.]+):?\s*(.*)$', value, flags=re.IGNORECASE)
        return (match.group(1), match.group(2).strip()) if match else (value, value)

    @classmethod
    def _tracker_names(cls, sheet, column):
        names = set()
        for row in sheet.iter_rows(min_row=3, values_only=True):
            value = cls._text(row[column - 1] if len(row) >= column else '')
            if value.startswith('Activity '):
                names.add(value)
                match = re.search(r'Activity\s+([0-9]+(?:\.[0-9]+)+)', value)
                if match: names.add(match.group(1))
        return names
